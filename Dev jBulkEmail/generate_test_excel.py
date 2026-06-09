import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

def create_testing_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "jBES Enterprise Testing Plan"

    # Define Headers
    headers = [
        "Test ID", "Phase", "Task Area", "Sub-Task / Condition", "QA Scenario & Validation", 
        "Expected Result", "Dependency / Integration Node", "Missing Link / Fallback Risk",
        "RAG Status", "Comments / Observations"
    ]
    
    # Styles
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    ws.append(headers)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Exhaustive Granular Test Data
    test_cases = [
        # 1. System Access & Auth (Phase 0)
        ("TC-F01.1", "Phase 0 (Foundation)", "Authentication", "First-Time Launch", "Launch app without config. Verify system locks access and prompts for Email to request OTP.", "UI locks on Auth Gate. OTP Dispatch triggered via Auth API.", "Auth Sync API", "No network access on first launch -> Blocks completely.", "Pending", ""),
        ("TC-F01.2", "Phase 0 (Foundation)", "Authentication", "Invalid OTP Submission", "Submit random 6-digit code to OTP validator.", "System rejects, throws error modal, increments failed attempt counter.", "Auth Sync API", "API Timeout -> UI freeze (Requires daemon thread).", "Pending", ""),
        ("TC-F01.3", "Phase 0 (Foundation)", "Authentication", "Terms & Conditions EULA", "Verify EULA modal blocks transmission UI until 'Accept' is clicked.", "Main UI disabled. Accepting writes 'eula_accepted: true' to config.", "UI State / Config JSON", "Corrupt config JSON -> Forces re-acceptance.", "Pending", ""),
        
        # 2. UX & Accessibility (Phase 0/1)
        ("TC-F02.1", "Phase 0 (Foundation)", "UX Accessibility", "Tab Sequence: Login Screen", "Navigate Login screen using only Tab key.", "Sequence: Email Input -> OTP Input -> Checkbox -> Verify Button.", "Tkinter UI focus", "Missing tabindex causes loop break.", "Pending", ""),
        ("TC-F02.2", "Phase 0 (Foundation)", "UX Accessibility", "Enter Key: Action Trigger", "Press <Enter> while focus is on Password field in Settings.", "System automatically triggers 'Save Account' function.", "Tkinter Event Binding", "No event mapping -> Manual mouse click required.", "Pending", ""),
        ("TC-F02.3", "Phase 0 (Foundation)", "UX Accessibility", "Mission Setup Safety Lock", "Press <Enter> while focus is in the 'Subject' input box.", "Focus shifts to Body textbox; does NOT trigger Mission Launch.", "Tkinter Event Binding", "Accidental launch misfire.", "Pending", ""),

        # 3. Transmission Engine & Dependencies (Phase 0/1)
        ("TC-D01.1", "Phase 1 (Development)", "Pre-Flight Validation", "Sender == Security Check", "Mismatch sender_email dropdown with active security_email.", "Launch aborted natively. UI throws mismatch warning.", "Core Engine Validation", "Bypass allows spoofing.", "Pending", ""),
        ("TC-D01.2", "Phase 1 (Development)", "Pre-Flight Validation", "Empty Subject/Body Check", "Attempt launch with empty body text.", "Launch aborted. UI warns 'Missing Content'.", "Core Engine Validation", "Sends blank emails.", "Pending", ""),
        ("TC-D02.1", "Phase 1 (Development)", "Network Resilience", "Mid-Loop SMTP Drop", "Disconnect WiFi exactly as `_smtp_send_with_retry` fires.", "Catches BrokenPipe/SMTPServerDisconnected. Retries twice. Falls back safely.", "smtplib exception handler", "Uncaught exception kills main thread.", "Pending", ""),
        ("TC-D02.2", "Phase 1 (Development)", "Throttling Defense", "10-Email Socket Rotation", "Track connection ID across 25 emails.", "Socket is explicitly quit and rotated using itertools.cycle exactly at email 10.", "itertools.cycle / smtplib", "Gmail Tarpit limits (Port blocked).", "Pending", ""),

        # 4. Heavy Attachments & Missing Links (Phase 1)
        ("TC-D03.1", "Phase 1 (Development)", "Payload Optimization", "Pre-encode Latency", "Load 3 attachments (5MB total). Verify encoding happens before loop starts.", "Log shows 'Build: <0.05s' per email in the loop.", "MIMEBase", "Encoding inside loop scales latency to O(N).", "Pending", ""),
        ("TC-D03.2", "Phase 1 (Development)", "Missing Link", "Attachment Deleted Mid-Flight", "Start mission, then quickly delete attached file from OS.", "FileNotFoundError caught. File skipped gracefully. Warning logged.", "FileSystem / OS", "Fatal crash if file strictly expected.", "Pending", ""),
        ("TC-D03.3", "Phase 1 (Development)", "Missing Link", "CSV File Deleted Mid-Flight", "Force delete active target CSV during active sending loop.", "CSV cache handles failure. Mission continues or pauses safely.", "FileSystem / pandas/csv", "Fatal crash on status update.", "Pending", ""),

        # 5. Checkpoints & State Recovery (Phase 0/1)
        ("TC-F04.1", "Phase 0 (Foundation)", "State Persistence", "Zero-Gap 50-item Flush", "Send 150 items. Verify disk I/O only spikes at 50, 100, 150.", "JSON checkpoint overwritten strictly at boundaries.", "JSON / Disk IO", "Memory leak if buffer not flushed.", "Pending", ""),
        ("TC-F04.2", "Phase 0 (Foundation)", "Missing Link", "Corrupted Mission JSON", "Manually edit Mission_X.json, invalidate JSON format, click Resume.", "JSONDecodeError caught. Modal drops corrupt file, UI survives.", "JSON module", "Recovery Modal crashes on load.", "Pending", ""),

        # 6. Advanced Relays & Closed-Loop (Phase 1)
        ("TC-D04.1", "Phase 1 (Development)", "Relay Engine", "SendGrid Payload Dispatch", "Post payload to SendGrid API with valid DPAPI key.", "HTTP 202 Accepted. Log shows API usage.", "requests / API", "Missing requests module.", "Pending", ""),
        ("TC-D04.2", "Phase 1 (Development)", "Relay Engine", "HTTP 429 Rate Limit Fallback", "Simulate HTTP 429 Too Many Requests response.", "Gracefully switch to SMTP `_smtp_send_with_retry` pool for remainder.", "HTTP Client", "Mission stalled permanently on API limit.", "Pending", ""),
        ("TC-D04.3", "Phase 1 (Development)", "Compression Engine", "Dynamic In-Memory Zip", "Load 2MB attachment. Verify compression bypasses disk.", "File zipped in RAM via zipfile module before HTTP POST.", "zipfile / io.BytesIO", "Disk I/O penalty if written to temp folder.", "Pending", ""),
        ("TC-D05.1", "Phase 1 (Development)", "Closed-Loop Feedback", "SMTPRecipientRefused", "Send to hard-bounced dummy email.", "Exception caught. Email instantly pushed to failed_list.", "smtplib / Closed-Loop", "Retries block queue on dead emails.", "Pending", ""),
        ("TC-D05.2", "Phase 1 (Development)", "Closed-Loop Feedback", "Failed CSV Generation", "Mission ends with 5 failed items. Click 'Retry Failed'.", "export_failed_csv writes failure log. Retry pulls exclusively from this new CSV.", "CSV Export", "Permission denied on export dir kills retry.", "Pending", "")
    ]

    for r_idx, row_data in enumerate(test_cases, 2):
        ws.append(row_data)
        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.alignment = wrap_align
            cell.border = thin_border
            if c_idx == 9: # RAG Status
                cell.alignment = center_align

    # Column Widths
    col_widths = {'A': 12, 'B': 22, 'C': 20, 'D': 30, 'E': 35, 'F': 35, 'G': 25, 'H': 30, 'I': 15, 'J': 30}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # RAG Status Data Validation
    dv = DataValidation(type="list", formula1='"Passed,Failed,Blocked,Pending"', allow_blank=True)
    dv.error = "Select a valid RAG status"
    dv.errorTitle = "Invalid Status"
    ws.add_data_validation(dv)
    
    rag_column = f"I2:I{len(test_cases) + 1}"
    dv.add(rag_column)

    # Conditional Formatting for RAG
    green_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    ws.conditional_formatting.add(rag_column, CellIsRule(operator='equal', formula=['"Passed"'], stopIfTrue=True, fill=green_fill))
    
    red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    ws.conditional_formatting.add(rag_column, CellIsRule(operator='equal', formula=['"Failed"'], stopIfTrue=True, fill=red_fill, font=Font(color="FFFFFF", bold=True)))
    
    amber_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    ws.conditional_formatting.add(rag_column, CellIsRule(operator='equal', formula=['"Blocked"'], stopIfTrue=True, fill=amber_fill))

    gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    ws.conditional_formatting.add(rag_column, CellIsRule(operator='equal', formula=['"Pending"'], stopIfTrue=True, fill=gray_fill))

    # Freeze top row
    ws.freeze_panes = "A2"

    wb.save("jBES_Enterprise_Testing_Plan_RAG.xlsx")
    print("Granular Task/Sub-Task Testing Excel with RAG status generated successfully.")

if __name__ == "__main__":
    create_testing_excel()
